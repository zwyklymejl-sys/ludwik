"""
BLOCK 5 - OUT-OF-SAMPLE BACKTEST

The patterns are calculated on the data known at the valuation period V.
If the input contains observations after V, we compare what each pattern
projected with what was actually paid afterwards.

For accident period i known at V:
    t_i  tenor reached at V (development at V + 1)
    C_i  cumulative paid at V
    R_i  best estimate at V (rbns + ibnr)
    U_i  = C_i + R_i  (ultimate as known at V - kept fixed for all later years)

A pattern pi projects the best estimate in the usual conditional way:
    expected paid in period V + h       = R_i * (pi_(t+h) - pi_(t+h-1)) / (1 - pi_t)
    expected cumulative paid at tenor s = C_i + R_i * (pi_s - pi_t) / (1 - pi_t)
(pi_s = 100% after the last tenor of the pattern).

Output (sheets of the comparison workbook):
  Backtest by tenor / by period - grouped views (written by excel_writer)
  Metrics - one row per aggregate x method x valuation: bias, SAE, MAE, SSE, RMSE,
            Pearson chi2, Huber loss, scaled versions, cumulative % errors, ranks
  Detail  - one row per aggregate x method x accident period x period after V
  Data    - the input data used (long format)

The three extra sheets are added by wrapping excel_writer.write_comparison_workbook
(see the end of this file), so only this module had to change.
"""
from dataclasses import dataclass, field
import math

import numpy as np

from . import excel_writer

HUBER_K = 1.345                # Huber constant (95% efficiency under normal errors)
MAD_TO_SIGMA = 0.6745          # MAD / 0.6745 = robust estimate of the standard deviation
HUBER_THRESHOLD = None         # fixed threshold in money; None = robust estimate per aggregate


@dataclass
class Backtest:
    valuation_label: str
    has_future: bool
    # by tenor
    tenors: np.ndarray             # 1..T
    actual_by_tenor: np.ndarray    # actual cumulative paid % of U (NaN where no cohort)
    count_by_tenor: np.ndarray
    expected_by_tenor: dict        # method name -> expected cumulative paid % of U
    # by period
    period_labels: list            # observation periods after V
    actual_by_period: np.ndarray
    expected_by_period: dict       # method name -> expected payments
    # detail rows and input data (for the Detail / Metrics / Data sheets)
    detail: list = field(default_factory=list)
    data_rows: list = field(default_factory=list)


def _pattern_value(pattern, tenor):
    """Cumulative pattern at a tenor (1-based); 100% after the last tenor."""
    if tenor <= len(pattern):
        return pattern[tenor - 1]
    return 1.0


def _label(text):
    """'2023' -> 2023 so that pivot tables sort numbers as numbers."""
    try:
        number = float(text)
        return int(number) if number.is_integer() else number
    except (TypeError, ValueError):
        return text


# ----------------------------------------------------------------------------
# Calculation
# ----------------------------------------------------------------------------
def run_backtest(full, known, valuation_pos, results):
    """full: all data; known: data cut at the valuation period; results: MethodResult list."""
    last_pos = full.latest_position
    n_known = known.n_accident
    has_future = last_pos > valuation_pos

    # ---- situation at the valuation period ------------------------------------
    tenor_at_v = valuation_pos - known.accident_position + 1            # t_i
    paid_at_v = np.nan_to_num(known.latest_diagonal(known.paid))         # C_i
    be_at_v = np.nan_to_num(known.latest_diagonal(known.best_estimate))  # R_i
    ultimate_at_v = paid_at_v + be_at_v                                  # U_i

    # full-data rows of the known accident periods (same order, they come first)
    full_paid = full.paid[:n_known, :]
    full_incremental = full.incremental_paid()[:n_known, :]

    # ---- by tenor ---------------------------------------------------------------
    n_tenors = full.n_development
    tenors = np.arange(1, n_tenors + 1)
    actual_sum = np.zeros(n_tenors)
    ultimate_sum = np.zeros(n_tenors)
    count = np.zeros(n_tenors, dtype=int)
    expected_sum = {r.method_name: np.zeros(n_tenors) for r in results}
    for i in range(n_known):
        for s in range(tenor_at_v[i] + 1, n_tenors + 1):
            k = s - 1                                                    # development index of tenor s
            if np.isnan(full_paid[i, k]) or ultimate_at_v[i] <= 0:
                continue
            actual_sum[s - 1] += full_paid[i, k]
            ultimate_sum[s - 1] += ultimate_at_v[i]
            count[s - 1] += 1
            for r in results:
                pi_t = _pattern_value(r.pattern, tenor_at_v[i])
                pi_s = _pattern_value(r.pattern, s)
                share = (pi_s - pi_t) / (1 - pi_t) if pi_t < 1 else 1.0
                expected_sum[r.method_name][s - 1] += paid_at_v[i] + be_at_v[i] * share

    actual_by_tenor = np.full(n_tenors, np.nan)
    expected_by_tenor = {name: np.full(n_tenors, np.nan) for name in expected_sum}
    for s in range(n_tenors):
        if count[s] > 0:
            actual_by_tenor[s] = actual_sum[s] / ultimate_sum[s]
            for name in expected_sum:
                expected_by_tenor[name][s] = expected_sum[name][s] / ultimate_sum[s]

    # ---- by observation period after V -------------------------------------------
    future_positions = list(range(valuation_pos + 1, last_pos + 1))
    period_labels = [full.period_labels[p] for p in future_positions]
    actual_by_period = np.zeros(len(future_positions))
    expected_by_period = {r.method_name: np.zeros(len(future_positions)) for r in results}
    for h, position in enumerate(future_positions, start=1):
        for i in range(n_known):
            k = position - known.accident_position[i]
            if k < full_incremental.shape[1] and not np.isnan(full_incremental[i, k]):
                actual_by_period[h - 1] += full_incremental[i, k]
            for r in results:
                pi_t = _pattern_value(r.pattern, tenor_at_v[i])
                if pi_t >= 1:
                    continue
                step = (_pattern_value(r.pattern, tenor_at_v[i] + h)
                        - _pattern_value(r.pattern, tenor_at_v[i] + h - 1))
                expected_by_period[r.method_name][h - 1] += be_at_v[i] * step / (1 - pi_t)

    # ---- detail: one row per method x accident period x period after V -------------
    valuation_label = known.period_labels[valuation_pos]
    detail = []
    for r in results:
        for i in range(n_known):
            t = int(tenor_at_v[i])
            pi_t = _pattern_value(r.pattern, t)
            for h, position in enumerate(future_positions, start=1):
                k = position - known.accident_position[i]
                if k >= full_paid.shape[1] or np.isnan(full_paid[i, k]):
                    continue                                             # no observation for this cell
                s = t + h
                pi_s = _pattern_value(r.pattern, s)
                pi_prev = _pattern_value(r.pattern, s - 1)
                if pi_t < 1:
                    expected_paid = be_at_v[i] * (pi_s - pi_prev) / (1 - pi_t)
                    expected_cum = paid_at_v[i] + be_at_v[i] * (pi_s - pi_t) / (1 - pi_t)
                else:
                    expected_paid = 0.0
                    expected_cum = paid_at_v[i] + be_at_v[i]
                actual_paid = float(np.nan_to_num(full_incremental[i, k]))
                actual_cum = float(full_paid[i, k])
                detail.append({
                    "aggregate": full.name, "method": r.method_code, "valuation": _label(valuation_label),
                    "AY": _label(known.accident_labels[i]), "period": _label(full.period_labels[position]),
                    "periods after valuation": h, "tenor at valuation": t, "tenor in period": s,
                    "paid at valuation": paid_at_v[i], "BE at valuation": be_at_v[i], "U at valuation": ultimate_at_v[i],
                    "pattern at valuation tenor": pi_t, "pattern at tenor in period": pi_s,
                    "expected paid in period": expected_paid, "actual paid in period": actual_paid,
                    "expected cum. paid": expected_cum, "actual cum. paid": actual_cum,
                })

    # ---- input data used (long format) ---------------------------------------------
    data_rows = []
    for i, accident in enumerate(full.accident_labels):
        for k in range(full.n_development):
            if np.isnan(full.paid[i, k]):
                continue
            position = full.accident_position[i] + k
            data_rows.append([full.name, _label(accident), _label(full.period_labels[position]), k,
                              float(full.paid[i, k]), float(full.rbns[i, k]), float(full.ibnr[i, k]),
                              float(full.best_estimate[i, k]), float(full.paid[i, k] + full.best_estimate[i, k]),
                              "yes" if position <= valuation_pos else "no"])

    return Backtest(valuation_label, has_future,
                    tenors, actual_by_tenor, count, expected_by_tenor,
                    period_labels, actual_by_period, expected_by_period,
                    detail, data_rows)


# ----------------------------------------------------------------------------
# Detail columns derived from the difference d = expected - actual
# ----------------------------------------------------------------------------
def huber_loss(d, threshold):
    """0.5 d^2 for |d| <= threshold, threshold (|d| - 0.5 threshold) above it."""
    if abs(d) <= threshold:
        return 0.5 * d * d
    return threshold * (abs(d) - 0.5 * threshold)


def huber_threshold(differences):
    """Fixed HUBER_THRESHOLD, or HUBER_K x robust sigma (MAD / 0.6745) of the differences.
    Falls back to HUBER_K x root mean square when the MAD is 0 (e.g. most differences equal)."""
    if HUBER_THRESHOLD is not None:
        return float(HUBER_THRESHOLD)
    if len(differences) == 0:
        return 0.0
    d = np.asarray(differences, dtype=float)
    mad = float(np.median(np.abs(d - np.median(d))))
    if mad > 0:
        return HUBER_K * mad / MAD_TO_SIGMA
    return HUBER_K * float(np.sqrt(np.mean(d * d)))


def _complete_detail(rows, threshold):
    for row in rows:
        d = row["expected paid in period"] - row["actual paid in period"]
        u = row["U at valuation"]
        row["difference paid (exp - act)"] = d
        row["abs difference"] = abs(d)
        row["squared difference"] = d * d
        row["Pearson term (d^2 / expected)"] = d * d / row["expected paid in period"] if row["expected paid in period"] > 0 else 0.0
        row["Huber loss"] = huber_loss(d, threshold)
        row["Huber threshold"] = threshold
        row["expected cum. % of U"] = row["expected cum. paid"] / u if u > 0 else None
        row["actual cum. % of U"] = row["actual cum. paid"] / u if u > 0 else None
        row["difference cum. % (exp - act)"] = (row["expected cum. % of U"] - row["actual cum. % of U"]) if u > 0 else None
        row["abs cum. difference"] = abs(row["expected cum. paid"] - row["actual cum. paid"])


# ----------------------------------------------------------------------------
# Metrics: one row per aggregate x method x valuation
# ----------------------------------------------------------------------------
METRIC_COLUMNS = ["n (AY x period)", "sum expected", "sum actual", "bias (exp - act)", "sum BE at valuation",
                  "|bias| / BE", "SAE (sum |diff|)", "MAE", "SSE (sum diff^2)", "RMSE", "Pearson chi2",
                  "Huber threshold", "Huber loss", "SAE / BE", "cum. MAE % of U", "cum. RMSE % pts"]
RANKED = ["|bias| / BE", "SAE (sum |diff|)", "RMSE", "Pearson chi2", "Huber loss", "cum. MAE % of U"]


def _metrics(rows):
    n = len(rows)
    sum_e = sum(r["expected paid in period"] for r in rows)
    sum_a = sum(r["actual paid in period"] for r in rows)
    be = sum(r["BE at valuation"] for r in rows if r["periods after valuation"] == 1)
    sae = sum(r["abs difference"] for r in rows)
    sse = sum(r["squared difference"] for r in rows)
    with_u = [r for r in rows if r["U at valuation"] > 0]
    sum_u = sum(r["U at valuation"] for r in with_u)
    m = {
        "n (AY x period)": n, "sum expected": sum_e, "sum actual": sum_a, "bias (exp - act)": sum_e - sum_a,
        "sum BE at valuation": be, "|bias| / BE": abs(sum_e - sum_a) / be if be > 0 else None,
        "SAE (sum |diff|)": sae, "MAE": sae / n if n else None, "SSE (sum diff^2)": sse,
        "RMSE": math.sqrt(sse / n) if n else None,
        "Pearson chi2": sum(r["Pearson term (d^2 / expected)"] for r in rows),
        "Huber threshold": rows[0]["Huber threshold"] if rows else None,
        "Huber loss": sum(r["Huber loss"] for r in rows),
        "SAE / BE": sae / be if be > 0 else None,
        "cum. MAE % of U": sum(r["abs cum. difference"] for r in with_u) / sum_u if sum_u > 0 else None,
        "cum. RMSE % pts": (math.sqrt(sum(r["difference cum. % (exp - act)"] ** 2 for r in with_u) / len(with_u))
                            if with_u else None),
    }
    return m


def build_metrics(backtests):
    """Complete the detail rows and return (detail rows, metric rows) over all aggregates."""
    all_detail, metric_rows = [], []
    for aggregate, bt in backtests.items():
        if not bt.detail:
            continue
        threshold = huber_threshold([r["expected paid in period"] - r["actual paid in period"] for r in bt.detail])
        _complete_detail(bt.detail, threshold)
        all_detail.extend(bt.detail)
        methods = []
        for r in bt.detail:
            if r["method"] not in methods:
                methods.append(r["method"])
        group = []
        for method in methods:
            m = _metrics([r for r in bt.detail if r["method"] == method])
            m.update({"aggregate": aggregate, "method": method, "valuation": _label(bt.valuation_label)})
            group.append(m)
        for name in RANKED:                                   # 1 = best (lowest) within the aggregate
            for m in group:
                value = m[name]
                m["rank " + name] = (None if value is None else
                                     1 + sum(1 for o in group if o[name] is not None and o[name] < value))
        metric_rows.extend(group)
    return all_detail, metric_rows


# ----------------------------------------------------------------------------
# Excel: Metrics, Detail and Data sheets appended to the comparison workbook
# ----------------------------------------------------------------------------
DETAIL_COLUMNS = ["aggregate", "method", "valuation", "AY", "period", "periods after valuation", "tenor at valuation",
                  "tenor in period", "paid at valuation", "BE at valuation", "U at valuation",
                  "pattern at valuation tenor", "pattern at tenor in period", "expected paid in period",
                  "actual paid in period", "difference paid (exp - act)", "abs difference", "squared difference",
                  "Pearson term (d^2 / expected)", "Huber threshold", "Huber loss", "expected cum. paid",
                  "actual cum. paid", "abs cum. difference", "expected cum. % of U", "actual cum. % of U",
                  "difference cum. % (exp - act)"]
DATA_COLUMNS = ["aggregate", "accident", "observation", "development", "paid", "rbns", "ibnr", "best estimate",
                "ultimate", "known at valuation"]
PERCENT_COLUMNS = {"pattern at valuation tenor", "pattern at tenor in period", "expected cum. % of U",
                   "actual cum. % of U", "difference cum. % (exp - act)", "|bias| / BE", "SAE / BE",
                   "cum. MAE % of U", "cum. RMSE % pts"}
INTEGER_COLUMNS = {"periods after valuation", "tenor at valuation", "tenor in period", "development",
                   "n (AY x period)"} | {"rank " + name for name in RANKED}


def _append_sheets(path, backtests):
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    detail, metric_rows = build_metrics(backtests)
    book = load_workbook(path)
    bold = Font(name="Arial", size=10, bold=True)
    plain = Font(name="Arial", size=10)

    def write_sheet(title, columns, rows):
        sheet = book.create_sheet(title)
        for j, name in enumerate(columns, start=1):
            sheet.cell(1, j, name).font = bold
            sheet.column_dimensions[get_column_letter(j)].width = max(11, min(len(name) + 2, 28))
        for i, row in enumerate(rows, start=2):
            for j, name in enumerate(columns, start=1):
                value = row[j - 1] if isinstance(row, list) else row.get(name)
                if value is None or (isinstance(value, float) and math.isnan(value)):
                    continue
                cell = sheet.cell(i, j, float(value) if isinstance(value, np.floating) else value)
                cell.font = plain
                if name in PERCENT_COLUMNS:
                    cell.number_format = "0.00%"
                elif name in INTEGER_COLUMNS:
                    cell.number_format = "0"
                elif isinstance(value, (float, np.floating)):
                    cell.number_format = "#,##0"
        sheet.freeze_panes = "B2"
        if rows:
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"

    metric_columns = ["aggregate", "method", "valuation"] + METRIC_COLUMNS + ["rank " + n for n in RANKED]
    write_sheet("Metrics", metric_columns, metric_rows)
    write_sheet("Detail", DETAIL_COLUMNS, detail)
    data_rows = []
    for bt in backtests.values():
        data_rows.extend(bt.data_rows)
    write_sheet("Data", DATA_COLUMNS, data_rows)
    book.save(path)


_write_comparison_workbook = excel_writer.write_comparison_workbook


def write_comparison_with_backtest(path, results_by_aggregate, backtests):
    """Comparison workbook as before + Metrics, Detail and Data sheets."""
    _write_comparison_workbook(path, results_by_aggregate, backtests)
    _append_sheets(path, backtests)


# run_pipeline imports this module before excel_writer, so it picks up the extended writer
excel_writer.write_comparison_workbook = write_comparison_with_backtest
